"""Multi-Format Document Ingestion for they.audio Podcast Platform.

Extracts structured paragraphs and titles from:
- PDF (.pdf)
- Microsoft Word (.docx)
- Markdown & Plain Text (.md, .txt)
- Direct Pasted Content
"""
from __future__ import annotations

import io
import os
import re
from typing import List, Optional, Tuple


def clean_line_breaks(text: str) -> str:
    """Fixes hyphenated word wraps and redundant line breaks from PDF/OCR."""
    # Fix hyphenation: "acous- \n tic" -> "acoustic"
    text = re.sub(r"(\w+)-\s*\n\s*(\w+)", r"\1\2", text)
    # Replace single line breaks inside paragraphs with space
    text = re.sub(r"(?<!\n)\n(?!\n)", " ", text)
    # Normalize multiple blank lines to double newlines
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def extract_from_pdf(content: bytes, fallback_title: str) -> Tuple[str, List[str]]:
    """Extracts title and paragraphs from PDF binary bytes."""
    import pypdf

    reader = pypdf.PdfReader(io.BytesIO(content))
    extracted_chunks: List[str] = []

    for page in reader.pages:
        txt = page.extract_text() or ""
        txt = clean_line_breaks(txt)
        if txt.strip():
            extracted_chunks.append(txt)

    full_text = "\n\n".join(extracted_chunks)
    paragraphs = [p.strip() for p in full_text.split("\n\n") if len(p.strip()) > 10]

    title = fallback_title
    if paragraphs:
        first_line = paragraphs[0].split("\n")[0].strip()
        if len(first_line) < 120:
            title = first_line
            # If the first paragraph was solely the title, omit from body
            if len(paragraphs[0]) == len(first_line):
                paragraphs = paragraphs[1:]

    return title, paragraphs


def extract_from_docx(content: bytes, fallback_title: str) -> Tuple[str, List[str]]:
    """Extracts title and paragraphs from Microsoft Word (.docx) bytes."""
    import docx

    doc = docx.Document(io.BytesIO(content))
    paragraphs: List[str] = []
    title = fallback_title

    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        if p.style and p.style.name.startswith("Heading 1") and title == fallback_title:
            title = txt
        else:
            paragraphs.append(txt)

    # Also capture table contents if any
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(f"Table row: {row_text}")

    if not paragraphs and title != fallback_title:
        paragraphs = [title]

    return title, paragraphs


def extract_from_text(content: str, fallback_title: str) -> Tuple[str, List[str]]:
    """Extracts title and paragraphs from plain text or Markdown."""
    cleaned = clean_line_breaks(content)
    raw_paras = [p.strip() for p in cleaned.split("\n\n") if p.strip()]

    title = fallback_title
    paragraphs: List[str] = []

    for idx, p in enumerate(raw_paras):
        # Detect markdown # Title
        if p.startswith("# ") and title == fallback_title:
            title = p[2:].strip()
            continue
        # Clean subheadings
        p_clean = re.sub(r"^#+\s*", "", p)
        if p_clean:
            paragraphs.append(p_clean)

    if not paragraphs and title != fallback_title:
        paragraphs = [title]

    return title, paragraphs


def extract_document(filename: str, content: bytes) -> Tuple[str, List[str]]:
    """Dispatches binary content to appropriate document parser based on file extension."""
    ext = os.path.splitext(filename)[1].lower()
    base_name = os.path.splitext(filename)[0].replace("-", " ").replace("_", " ").title()

    if ext == ".pdf":
        return extract_from_pdf(content, fallback_title=base_name)
    elif ext in (".docx", ".doc"):
        return extract_from_docx(content, fallback_title=base_name)
    elif ext in (".txt", ".md", ".markdown", ".rst"):
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            text = content.decode("latin-1", errors="ignore")
        return extract_from_text(text, fallback_title=base_name)
    elif ext in (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"):
        raise ValueError(
            f"Image screenshot '{filename}' requires OCR processing. "
            "Please upload via the they.audio browser workstation which runs neural WebAssembly OCR, "
            "or send recognized text to /api/extract/text."
        )
    else:
        # Fallback text decoder
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            text = content.decode("latin-1", errors="ignore")
        return extract_from_text(text, fallback_title=base_name)
