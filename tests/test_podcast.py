"""Unit and regression tests for Document Extraction and Podcast Cadence Engine."""
from __future__ import annotations

import io
import pytest
import docx
from fastapi.testclient import TestClient

from app.main import app
from app.extract.document import extract_document
from app.podcast.cadence import analyze_sentence_cadence, shape_article_into_podcast


def test_cadence_sentence_analysis() -> None:
    # 1. Question sentence
    disp, spch, pause_ms, mood = analyze_sentence_cadence("Why does human pacing matter?", 1)
    assert mood == "question"
    assert pause_ms >= 500

    # 2. Revelation sentence
    disp2, spch2, pause_ms2, mood2 = analyze_sentence_cadence("The truth is, without pauses speech feels robotic.", 2)
    assert mood2 == "revelation"
    assert "<break" in spch2

    # 3. Em-dash and colon pauses
    disp3, spch3, _, _ = analyze_sentence_cadence("Here is the key: audiobooks need cadence — natural breathing.", 3)
    assert "<break time=" in spch3


def test_cadence_punchy_and_enumeration() -> None:
    # 1. Enumeration sentence
    disp, spch, pause_ms, mood = analyze_sentence_cadence("First, let us examine the acoustic foundation.", 1)
    assert mood == "enumeration"
    assert pause_ms >= 300

    # 2. Punchy takeaway
    disp2, spch2, pause_ms2, mood2 = analyze_sentence_cadence("It worked.", 2)
    assert mood2 == "punchy"
    assert pause_ms2 >= 450

    # 3. Parenthetical aside
    disp3, spch3, _, _ = analyze_sentence_cadence("The soundscape (which is procedural) brings warmth.", 3)
    assert "— which is procedural —" in disp3


def test_document_extract_docx() -> None:
    doc = docx.Document()
    doc.add_heading("Architecture of Sound", level=1)
    doc.add_paragraph("First paragraph discussing acoustic design.")
    doc.add_paragraph("Second paragraph explaining frequency analysis.")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    title, paras = extract_document("architecture.docx", buf.read())
    assert title == "Architecture of Sound"
    assert len(paras) == 2


def test_document_extract_pdf_fallback() -> None:
    import pypdf
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    buf.seek(0)

    title, paras = extract_document("test_doc.pdf", buf.read())
    assert "Test Doc" in title
    assert isinstance(paras, list)


def test_api_extract_text_podcast_mode() -> None:
    client = TestClient(app)
    text = (
        "# Why Podcasting Works\n\n"
        "Podcasting is about intimacy and pacing. Why does a voice sound real?\n\n"
        "In fact, human speech pauses. Think about it: without silence, words lose meaning.\n\n"
        "First, cadence sets the rhythm. It worked."
    )
    res = client.post("/api/extract/text", json={
        "text": text,
        "title": "Why Podcasting Works",
        "podcast_mode": True,
    })
    assert res.status_code == 200
    data = res.json()
    assert "Why Podcasting Works" in data["title"]
    assert len(data["sentences"]) >= 5
    assert any(s["mood"] == "question" for s in data["sentences"])
    assert any(s["mood"] == "revelation" for s in data["sentences"])
    assert any(s["mood"] == "punchy" for s in data["sentences"])


def test_api_extract_file_markdown() -> None:
    client = TestClient(app)
    file_bytes = io.BytesIO(b"# Monograph on Taste\n\nParagraph one.\n\nParagraph two.")
    res = client.post(
        "/api/extract/file",
        files={"file": ("taste.md", file_bytes, "text/markdown")},
        data={"podcast_mode": "true"},
    )
    assert res.status_code == 200
    data = res.json()
    assert "Monograph on Taste" in data["title"]
    assert len(data["sentences"]) >= 2


def test_extract_document_rejects_raw_image() -> None:
    with pytest.raises(ValueError, match="requires OCR processing"):
        extract_document("screenshot.png", b"\x89PNG\r\n\x1a\n...")


def test_dialogue_strips_cheesy_ai_banter() -> None:
    from app.podcast.dialogue import strip_cheesy_banter

    raw = (
        "Welcome back to the deep dive! Buckle up! "
        "Wow, that's so fascinating! That's a great point! "
        "Cognitive load theory explains how working memory processes new information."
    )
    cleaned = strip_cheesy_banter(raw)
    assert "deep dive" not in cleaned.lower()
    assert "buckle up" not in cleaned.lower()
    assert "fascinating" not in cleaned.lower()
    assert "great point" not in cleaned.lower()
    assert "Cognitive load theory explains" in cleaned


def test_dialogue_avoids_ping_pong_and_injects_cognitive_focus() -> None:
    from app.podcast.dialogue import shape_article_into_dialogue

    paragraphs = [
        "Working memory has a strict bottleneck when decoding unfamiliar technical abstractions, which causes readers to lose focus after ninety seconds of unbroken prose.",
        "However, when a second perspective interrupts the monologue to ask a concrete question, the brain resets its attention loop and anchors the concept in long-term memory.",
        "In fact, active retrieval matters far more than passive rereading: testing yourself on the core mechanism forces neural consolidation.",
    ]
    sents = shape_article_into_dialogue(
        title="How the Brain Locks In Complex Ideas",
        paragraphs=paragraphs,
        voice_a="en-US-AndrewMultilingualNeural",
        voice_b="en-US-AvaMultilingualNeural",
    )
    assert len(sents) > len(paragraphs) + 1
    speakers = [s["speaker"] for s in sents]
    assert "A" in speakers
    assert "B" in speakers

    # Ensure voices are properly bound to speakers
    for s in sents:
        if s["speaker"] == "A":
            assert s["voice_id"] == "en-US-AndrewMultilingualNeural"
            assert s["speaker_name"] == "Anchor"
        else:
            assert s["voice_id"] == "en-US-AvaMultilingualNeural"
            assert s["speaker_name"] == "Focus Partner"

    # Anti-Ping-Pong check: Speaker B must provide cognitive focus/clarification turns,
    # not just read raw source paragraphs verbatim in a 1-to-1 alternation
    b_texts = " ".join(s["display_text"].lower() for s in sents if s["speaker"] == "B")
    assert any(
        cue in b_texts
        for cue in ["hold on", "wait", "plain", "so if", "why", "lock", "break", "mean", "takeaway", "unpack"]
    )
    # Anti-Banter check
    all_text = " ".join(s["display_text"].lower() for s in sents)
    assert "welcome back" not in all_text
    assert "buckle up" not in all_text


def test_api_extract_text_dialogue_mode() -> None:
    client = TestClient(app)
    text = (
        "# Distributed Consensus in Plain Sight\n\n"
        "Byzantine fault tolerance requires nodes to agree on state even when some actors transmit conflicting messages across an unreliable network.\n\n"
        "However, the real bottleneck is not computation — it is the quadratic message complexity required for every round of voting."
    )
    res = client.post("/api/extract/text", json={
        "text": text,
        "title": "Distributed Consensus in Plain Sight",
        "podcast_mode": True,
        "dialogue_mode": True,
        "voice_a": "en-US-AndrewMultilingualNeural",
        "voice_b": "en-US-AvaMultilingualNeural",
    })
    assert res.status_code == 200
    data = res.json()
    assert data["extractor_used"] == "dialogue_text"
    speakers = {s["speaker"] for s in data["sentences"]}
    assert speakers == {"A", "B"}
    voices = {s["voice_id"] for s in data["sentences"]}
    assert "en-US-AndrewMultilingualNeural" in voices
    assert "en-US-AvaMultilingualNeural" in voices


def test_api_extract_file_dialogue_mode() -> None:
    client = TestClient(app)
    file_bytes = io.BytesIO(
        b"# Cognitive Focus Architecture\n\n"
        b"Dense technical monographs often overwhelm working memory because clauses stack without a pause for synthesis.\n\n"
        b"Crucially, breaking a thesis into an anchor explanation and a clarifying question keeps the listener locked in."
    )
    res = client.post(
        "/api/extract/file",
        files={"file": ("focus.md", file_bytes, "text/markdown")},
        data={
            "podcast_mode": "true",
            "dialogue_mode": "true",
            "voice_a": "en-GB-RyanNeural",
            "voice_b": "en-GB-SoniaNeural",
        },
    )
    assert res.status_code == 200
    data = res.json()
    assert data["extractor_used"] == "dialogue_document"
    speakers = {s["speaker"] for s in data["sentences"]}
    assert speakers == {"A", "B"}
    voices = {s["voice_id"] for s in data["sentences"]}
    assert "en-GB-RyanNeural" in voices
    assert "en-GB-SoniaNeural" in voices

